import os
import re
import traceback
from datetime import date

import openpyxl
import pandas as pd
import streamlit as st

import assumptions_parser
import config_loader
import kardin_parser
import leasing_parser
import recoveries_parser
import expense_parser
import capex_parser
import forecast_parser
import xtra_parser
import tracker_writer

st.set_page_config(page_title="Kardin Budget QC", layout="wide")

st.title("Kardin Budget QC")
st.caption("All 7 file buckets. Cross-bucket checks (4/5/6 → 1, 7 → 2) run automatically "
           "when the earlier bucket has already been analyzed this session for the same building.")

if 'building' not in st.session_state:
    st.session_state.building = ''
if 'bucket_results' not in st.session_state:
    st.session_state.bucket_results = {}  # (bucket_num, building) -> dict


STATUS_ICON = {'Passed': '✅', 'Flagged': '⚠️', 'Skipped': '⏭️'}


def show_checklist(checklist):
    """Named rules + pass/flag/skip status for this bucket's run - same idea
    as ga-automation's QC engine (CHECK_1..CHECK_7 with a status each)
    instead of just a flat pile of findings, so it's immediately clear which
    specific checks ran clean, which flagged something, and which couldn't
    run at all because a required file/cross-bucket input was missing."""
    if not checklist:
        return
    st.subheader("QC Checklist")
    for c in checklist:
        icon = STATUS_ICON.get(c['status'], '•')
        label = f"{icon} **{c['name']}** - {c['status']}"
        if c['count']:
            label += f" ({c['count']})"
        st.markdown(label)


def show_expense_categorization(building_key, cost_center_code=None):
    """Cross-bucket enrichment shown on Bucket 1: how much of the budget is
    Contingency / Contract (vendor-backed) / Misc-Other / no vendor note on
    file, using Bucket 4's (Expense Detail) vendor-level line items - the
    only place that detail exists. Silently does nothing if bucket 4 hasn't
    been analyzed yet for this building (lookup_bucket also falls back to an
    "(All Buildings)" run of bucket 4 if that's how it was run)."""
    b4_entry = lookup_bucket(4, building_key)
    if not b4_entry:
        return
    line_rows = b4_entry['results'].get('line_rows')
    if not line_rows:
        return
    source_label = b4_entry.get('expense_detail_name', 'Expense Detail')
    cat_findings = kardin_parser.check_expense_categorization(
        line_rows, source_label=source_label, cost_center_filter=cost_center_code)
    if not cat_findings:
        return
    st.subheader("Budget Categorization (from Expense Detail)")
    st.caption("Best-effort read on how much of the budget is Contingency, vendor-backed (Contract), "
               "tagged Misc/Other by GL account, or has no vendor/note on file - a heuristic to help "
               "review, not a rule violation.")
    for f in cat_findings:
        st.markdown(f"- {f['Comment']}")


def _file_label(f):
    """Display name for either a Streamlit UploadedFile or a committed-file
    path string - both are valid inputs to load_or_upload_file's callers."""
    return getattr(f, 'name', None) or os.path.basename(str(f))


def _seek0(f):
    """Rewind a file-like object before re-reading it a second time in the
    same run - a no-op for a committed-file path string, which has no
    seek() and doesn't need one."""
    if hasattr(f, 'seek'):
        f.seek(0)


def load_or_upload_file(assumption_key, label, file_type, key_prefix):
    """Returns a committed file (path string, from this property's
    config.yaml assumption_files map) if one exists, else renders a
    required uploader - or, when a committed file IS found, renders an
    OPTIONAL override uploader instead (an upload always wins over the
    committed file for that run, without needing to edit the config).
    Returns None if nothing is available yet. Works transparently with
    downstream code (openpyxl.load_workbook, parse_expense_detail, etc.)
    since both a path string and an UploadedFile are valid inputs to those -
    just use _seek0()/_file_label() instead of .seek()/.name directly."""
    slug = st.session_state.get('property_slug')
    cfg = st.session_state.get('property_cfg') or {}
    filename = (cfg.get('assumption_files') or {}).get(assumption_key)
    committed = config_loader.committed_file_path(slug, filename)
    if committed:
        st.caption(f"✓ Using committed {label}: `{os.path.basename(committed)}` (from this property's config)")
        override = st.file_uploader(f"Override {label} (optional)", type=file_type,
                                    key=f'{key_prefix}_{assumption_key}_override')
        return override if override else committed
    return st.file_uploader(label, type=file_type, key=f'{key_prefix}_{assumption_key}_upload')


def show_assumptions_check(monthly_rows, key_prefix):
    """Optional cross-check: GRP's own GL-level monthly budget assumptions
    (a separate planning workbook GRP provides, not read from Kardin) vs
    Kardin's actual Monthly Budget Detail for this building. Nothing here
    is required for the rest of the tool - upload is entirely optional."""
    with st.expander("Cross-check against GRP Budget Assumptions (optional)"):
        st.caption("Committed per-property in config.yaml if set up; otherwise upload GRP's own planning "
                   "workbook (e.g. '2027 GRP Budget Assumptions.xlsx') to check whether Kardin's Monthly "
                   "Budget Detail actually reflects what GRP assumed, GL by GL.")
        xlsx_file = load_or_upload_file('grp_budget_assumptions', "GRP Budget Assumptions (.xlsx)", "xlsx", key_prefix)
        if not xlsx_file:
            return
        try:
            wb = openpyxl.load_workbook(xlsx_file, data_only=True, read_only=True)
            sheet_names = wb.sheetnames
        except Exception:
            st.error("Couldn't read that file as an Excel workbook.")
            return
        col1, col2 = st.columns(2)
        with col1:
            sheet_name = st.selectbox("Which tab is this building?", sheet_names, key=f'{key_prefix}_assump_sheet')
        with col2:
            budget_year = st.number_input(
                "Budget year to compare (the 'next year' column set - the workbook usually also carries an "
                "earlier reforecast tail that should NOT be compared against this Kardin file)",
                min_value=2000, max_value=2100, value=date.today().year + 1, key=f'{key_prefix}_assump_year')
        try:
            _seek0(xlsx_file)
            assumption_rows = assumptions_parser.parse_grp_budget_assumptions(xlsx_file, sheet_name)
            findings = assumptions_parser.check_assumptions_vs_bucket1(
                assumption_rows, monthly_rows, source_label=f"{_file_label(xlsx_file)} ({sheet_name})",
                budget_year=int(budget_year))
        except Exception:
            report_error()
            return
        if not findings:
            st.info("No mismatches found against this assumption sheet for this budget year.")
            return
        for f in findings:
            st.markdown(f"- {f['Comment']}")


def show_interest_check(monthly_rows, key_prefix):
    """Optional, PURELY INFORMATIONAL cross-check: GRP's construction-loan
    Interest Carry schedule vs Kardin's GL 801110 (Interest Expense) for
    this building. Never asserts a mismatch is wrong - see
    assumptions_parser.check_interest_vs_bucket1's docstring for why."""
    with st.expander("Cross-check against Interest Calculations (optional, informational only)"):
        st.caption("Upload GRP's own Interest Calculations workbook to compare its construction-loan "
                   "interest schedule against Kardin's GL 801110 for this building. Purely informational - "
                   "the two can legitimately differ (refinancing, a different note structure, timing) and "
                   "this check can't tell which; it's for you to judge, not a rule violation.")
        xlsx_file = load_or_upload_file('interest_calculations', "Interest Calculations (.xlsx)", "xlsx", key_prefix)
        if not xlsx_file:
            return
        try:
            wb = openpyxl.load_workbook(xlsx_file, data_only=True, read_only=True)
            sheet_names = wb.sheetnames
        except Exception:
            st.error("Couldn't read that file as an Excel workbook.")
            return
        col1, col2 = st.columns(2)
        with col1:
            sheet_name = st.selectbox("Which tab is this property?", sheet_names, key=f'{key_prefix}_interest_sheet')
        with col2:
            budget_year = st.number_input("Budget year to compare", min_value=2000, max_value=2100,
                                          value=date.today().year + 1, key=f'{key_prefix}_interest_year')
        try:
            _seek0(xlsx_file)
            sections = assumptions_parser.parse_interest_calculations(xlsx_file, sheet_name)
        except Exception:
            report_error()
            return
        if not sections:
            st.info("No '... Interest Calculation' section found on that tab.")
            return
        chosen = st.selectbox("Which section is THIS building?", [s['title'] for s in sections],
                              key=f'{key_prefix}_interest_section')
        section = next(s for s in sections if s['title'] == chosen)
        findings = assumptions_parser.check_interest_vs_bucket1(
            section, monthly_rows, int(budget_year), source_label=f"{_file_label(xlsx_file)} ({sheet_name})")
        if not findings:
            st.info("No difference beyond tolerance found for this building/year.")
            return
        for f in findings:
            st.markdown(f"- {f['Comment']}")


def show_leasing_assumptions_check(occupancy_rows, key_prefix):
    """Optional cross-check: GRP's own Leasing Assumptions workbook (vacant/
    expiring suites GRP told the PM to model - RSF, commencement date, term,
    free rent) vs whether Kardin's Occupancy Summary actually reflects it.
    Uses the current property's cost_centers config to map each building
    name in the workbook (e.g. '20 Riverside') to its Kardin cost center
    (e.g. 'west20') - without a property selected, building names can't be
    resolved and nothing will match."""
    with st.expander("Cross-check against Leasing Assumptions (optional)"):
        st.caption("Upload GRP's own Leasing Assumptions workbook (e.g. '2027 Leasing Assumptions_Final.xlsx') "
                   "to check whether Kardin's Occupancy Summary reflects the vacant/expiring suite assumptions "
                   "GRP gave the PM. The sheet is reused year over year with one stacked block per budget cycle - "
                   "only the block for the year you enter below is used.")
        xlsx_file = load_or_upload_file('leasing_assumptions', "Leasing Assumptions (.xlsx)", "xlsx", key_prefix)
        if not xlsx_file:
            return
        try:
            wb = openpyxl.load_workbook(xlsx_file, data_only=True, read_only=True)
            sheet_names = wb.sheetnames
        except Exception:
            st.error("Couldn't read that file as an Excel workbook.")
            return
        col1, col2 = st.columns(2)
        with col1:
            sheet_name = st.selectbox("Which tab is this property?", sheet_names, key=f'{key_prefix}_lease_assump_sheet')
        with col2:
            budget_year = st.number_input(
                "Budget year (picks the matching 'Leasing Assumptions - ... {year} Budget' block)",
                min_value=2000, max_value=2100, value=date.today().year + 1, key=f'{key_prefix}_lease_assump_year')
        property_cfg = st.session_state.get('property_cfg')
        building_to_cc = {cc['name']: cc['code'] for cc in (property_cfg or {}).get('cost_centers', [])
                          if cc.get('name') and cc.get('code')}
        if not building_to_cc:
            st.warning("No property selected above (or it has no cost centers configured) - building names in "
                       "this workbook (e.g. '20 Riverside') can't be mapped to a Kardin cost center, so nothing "
                       "will match. Select a Property above first.")
        try:
            _seek0(xlsx_file)
            assumption_rows = assumptions_parser.parse_leasing_assumptions(xlsx_file, sheet_name, int(budget_year))
            findings = assumptions_parser.check_leasing_assumptions_vs_bucket2(
                assumption_rows, occupancy_rows, building_to_cc, int(budget_year),
                source_label=f"{_file_label(xlsx_file)} ({sheet_name})")
        except Exception:
            report_error()
            return
        if not assumption_rows:
            st.info(f"No 'Leasing Assumptions - ... {int(budget_year)} Budget' block found on that tab.")
            return
        if not findings:
            st.info(f"No mismatches found against the {int(budget_year)} Budget block.")
            return
        for f in findings:
            st.markdown(f"- **[{f['Priority']}]** {f['Comment']}")


def show_capex_plan_check(capex_line_rows, key_prefix):
    """Optional cross-check: GRP's own 5-Year CapEx Plan's per-building
    annual total (building-improvement capital only - HVAC, roof, etc.) vs
    Kardin's actual Capex.pdf for the same building. Uses the current
    property's cost_centers config to map building names (e.g. '20
    Riverside') to Kardin cost centers (e.g. 'west20')."""
    with st.expander("Cross-check against 5-Year CapEx Plan (optional)"):
        st.caption("Upload GRP's own 5-Year CapEx Plan (e.g. 'GRP 5-Year CapEx Plan Riverside Labs....xlsx') "
                   "to check whether Kardin's Capex.pdf reflects the planned building-improvement capital "
                   "for the target year. Only GL 154500/171300 are compared - Leasing Commissions/Tenant "
                   "Improvements are a different capital category, already checked elsewhere in this bucket.")
        xlsx_file = load_or_upload_file('capex_plan', "5-Year CapEx Plan (.xlsx)", "xlsx", key_prefix)
        if not xlsx_file:
            return
        budget_year = st.number_input(
            "Budget year to compare", min_value=2000, max_value=2100,
            value=date.today().year + 1, key=f'{key_prefix}_capex_plan_year')
        property_cfg = st.session_state.get('property_cfg')
        building_to_cc = {cc['name']: cc['code'] for cc in (property_cfg or {}).get('cost_centers', [])
                          if cc.get('name') and cc.get('code')}
        if not building_to_cc:
            st.warning("No property selected above (or it has no cost centers configured) - building names "
                       "in this workbook can't be mapped to a Kardin cost center. Select a Property above first.")
        try:
            _seek0(xlsx_file)
            plan_totals = assumptions_parser.parse_capex_plan_portfolio_summary(xlsx_file, int(budget_year))
            findings = assumptions_parser.check_capex_plan_vs_bucket5(
                plan_totals, capex_line_rows, building_to_cc, int(budget_year), source_label=_file_label(xlsx_file))
        except Exception:
            report_error()
            return
        if not plan_totals:
            st.info(f"No {int(budget_year)} column found in that plan's Portfolio Summary tab.")
            return
        if not findings:
            st.info(f"Ties out against the {int(budget_year)} Portfolio Summary for every mapped building.")
            return
        for f in findings:
            st.markdown(f"- {f['Comment']}")


def show_expense_yoy_comparison(current_line_rows, key_prefix):
    """Optional: compare this year's Expense Detail against last year's -
    what changed GL by GL, whether Contingency/Contract/Misc/No-vendor-note
    category totals shifted, and which specific lines are driving the
    overall change (with the current year's own vendor-level detail
    attached to the top drivers, for the "why")."""
    with st.expander("Compare to Prior Year Expense Detail (optional)"):
        st.caption("Upload last year's Expense Detail (.pdf) for this same property to see what changed, "
                   "whether Contingency stayed flat, and what's driving any increase.")
        prior_pdf = load_or_upload_file('prior_year_expense_detail', "Prior Year Expense Detail (.pdf)", "pdf", key_prefix)
        if not prior_pdf:
            return
        property_cfg = st.session_state.get('property_cfg')
        cc_list = [cc for cc in (property_cfg or {}).get('cost_centers', []) if cc.get('name') and cc.get('code')]
        cc_options = ['(All buildings combined)'] + [cc['name'] for cc in cc_list]
        chosen_cc = st.selectbox("Compare for which building?", cc_options, key=f'{key_prefix}_yoy_cc')
        cost_center_code = None
        if chosen_cc != cc_options[0]:
            match = next((cc for cc in cc_list if cc['name'] == chosen_cc), None)
            cost_center_code = match['code'] if match else None

        try:
            prior_line_rows, _ = expense_parser.parse_expense_detail(prior_pdf)
        except Exception:
            report_error()
            return
        if not prior_line_rows:
            st.warning("Couldn't find any line items in that file - it may be an incomplete/broken export "
                       "(e.g. every page shows only a loading placeholder instead of real report content - "
                       "check the source PDF and re-export if so).")
            return

        cmp = kardin_parser.compare_expense_detail_yoy(current_line_rows, prior_line_rows,
                                                        cost_center_filter=cost_center_code)
        st.metric("Total", f"${cmp['total_current']:,.0f}",
                  delta=f"${cmp['total_delta']:+,.0f} ({cmp['total_pct_change']:+.1f}%)")

        st.subheader("By category")
        st.caption("Same Contingency/Contract/Misc/No-vendor-note rule as the Budget Categorization panel "
                   "on Bucket 1 - shows whether each shifted year over year.")
        cat_df = pd.DataFrame([{
            'Category': c['category'], 'Prior $': c['prior'], 'Current $': c['current'],
            'Delta $': c['delta'], 'Delta %': round(c['pct_change'], 1),
        } for c in cmp['category_comparison']])
        st.dataframe(cat_df, use_container_width=True, hide_index=True)

        st.subheader("Top drivers of the change")
        st.caption("Ranked by absolute $ delta - the GL lines doing the most to explain the total change.")
        top_gls = cmp['gl_comparison'][:15]
        gl_df = pd.DataFrame([{
            'GL': g['gl'], 'Label': g['label'], 'Prior $': g['prior'], 'Current $': g['current'],
            'Delta $': g['delta'], 'Delta %': round(g['pct_change'], 1),
            'Flag': 'NEW' if g['is_new'] else ('DROPPED' if g['is_dropped'] else ''),
        } for g in top_gls])
        st.dataframe(gl_df, use_container_width=True, hide_index=True)

        with st.expander("Line-item detail behind the top 5 drivers (this year)"):
            for g in top_gls[:5]:
                if not g['current_line_items']:
                    continue
                st.markdown(f"**{g['gl']} {g['label']}** (delta ${g['delta']:+,.0f})")
                for li in g['current_line_items']:
                    st.markdown(f"- {li['description'] or '(no description)'}: ${li['total']:,.0f}")


def show_stats(stats):
    st.json(stats, expanded=False)


def show_findings(findings):
    if not findings:
        st.info("No issues found by the current checks.")
        return
    df = pd.DataFrame(findings)[['Priority', 'Report Section', 'GL Acct', 'Line Item',
                                  'Budget Year', 'Comment', 'Source Check']]
    must_fix = df[df['Priority'] == 'Must Fix']
    discuss = df[df['Priority'] != 'Must Fix']
    st.subheader(f"Must Fix ({len(must_fix)})")
    st.caption("Straightforward, rule-based - SOP thresholds, tie-outs, cross-report mismatches.")
    st.dataframe(must_fix, use_container_width=True, hide_index=True)
    st.subheader(f"For Discussion ({len(discuss)})")
    st.caption("Flagged by the parser but needs a human read before it's a real comment.")
    st.dataframe(discuss, use_container_width=True, hide_index=True)


def run_button(key, requirements):
    """requirements: {label: is_satisfied_bool}. Missing files no longer
    block the run - they flag it instead (each bucket's run() adds its own
    "Missing file" finding for whatever wasn't provided) so partial uploads
    still produce whatever findings ARE possible from what's there."""
    missing = [label for label, ok in requirements.items() if not ok]
    if missing:
        st.caption(f"⚠️ Missing: {', '.join(missing)}. Findings that depend on these will be "
                   "flagged, not skip the whole run - click Run Analysis to see what's available now.")
    return st.button("Run Analysis", type="primary", key=key)


def report_error():
    st.error("Parsing failed. Full traceback below - this is exactly the kind of bug we want to surface fast.")
    st.code(traceback.format_exc())


# Riverside Labs' own bucket-2 file was "...Rent-Lab-Mnthly.pdf" - not
# spelled out. Every rule that keys off "monthly" (to find/exclude a Monthly
# Detail file) checks all of these instead of just the one literal spelling.
MONTHLY_TOKENS = ['monthly', 'mnthly', 'mthly']

# Buckets 2/3/4/5/7 don't always get per-building files the way Bucket 1/6
# do at Lexington Labs (B<n> tags) - some PMs export these portfolio-wide,
# covering every building in one file (e.g. one combined Recovery Calc Est
# instead of five). Forcing a single specific "Building name" to run one of
# those doesn't fit, so this sentinel lets a bucket be marked as applying to
# every building at once instead of picking one.
ALL_BUILDINGS = '(All Buildings)'


def lookup_bucket(n, building):
    """Cross-bucket lookup: prefer this specific building's results, fall
    back to an ALL_BUILDINGS run of that bucket if this building doesn't have
    its own (e.g. Bucket 2 was run portfolio-wide but Bucket 7 is being run
    for one specific building)."""
    return (st.session_state.bucket_results.get((n, building))
            or st.session_state.bucket_results.get((n, ALL_BUILDINGS)))


def building_selector(prefix, building):
    """Renders the 'this bucket covers all buildings' checkbox for a
    single-file bucket. Returns (effective_building_key, requirement_met) -
    use effective_building_key wherever the bucket's OWN results get stored/
    read; keep using the plain `building` value for any cross-bucket lookup
    into a per-building bucket like 1 or 2, since those still need one
    specific building's data even when THIS bucket is portfolio-wide."""
    all_buildings = st.checkbox(
        "This bucket's files cover ALL buildings (one portfolio-wide export, not building-specific)",
        key=f'{prefix}_all_buildings',
    )
    if all_buildings:
        return ALL_BUILDINGS, True
    return building, bool(building)


def has_bucket_tag(fname_lower, n):
    """True if the filename carries an explicit 'B1'/'B2'/... bucket-number
    tag (as a whole token, so 'b1' doesn't match inside 'b10' etc). Some
    properties' files are pre-labeled this way (discovered on Lexington
    Labs), which is a much stronger signal than keyword-guessing alone."""
    return re.search(rf'(?<![a-z0-9])b0*{n}(?![a-z0-9])', fname_lower) is not None


def exclude_3way_reports(files):
    """Bucket 6's 2-way ('2025B v 2025F') Detail/Monthly Detail files can
    collide with Bucket 1's 3-way ('2025B v 2025F v 2026B') Detail files
    when both carry the SAME B<n> building tag and match the same generic
    'detail' keyword rule - confirmed against real Lex Labs data run through
    the full shared file pool (the normal way this app is used): bucket 6's
    classification silently picked bucket 1's own Detail/Monthly Detail
    file instead of its own, for every one of 5 buildings, producing 0
    findings instead of real ones (bucket 1's file structurally doesn't
    match bucket 6's 2-way parser, so it just quietly extracts nothing).
    A 3-way report's filename always has TWO ' v ' separators (year v year
    v year); a 2-way report has only one - filter out anything with 2+
    before bucket 6 (or anything else that's specifically 2-way) ever sees
    it, so the B<n>-tag+keyword match can't cross-pick the wrong bucket's
    file even though both are tagged for the same building."""
    return [f for f in (files or []) if f.name.lower().count(' v ') < 2]


def classify_and_pick(files, slot_rules, key_prefix, bucket_number=None, building_hint=None):
    """
    All files for every bucket are dropped into ONE global uploader; this
    guesses which uploaded file goes in which slot for THIS bucket, then
    renders an editable dropdown per slot (defaulted to the guess, but
    listing every uploaded file) so a naming-convention mismatch never
    blocks you - worst case you just pick manually.

    Matching is up to three-pass: first among files carrying an explicit
    'B{bucket_number}' tag (when bucket_number is given); then, when
    building_hint is given (the current Building name), among files whose
    name contains it as a substring - properties with multiple buildings
    but NO B<n> tags (e.g. Riverside Labs' "20 Riverside" vs "1 Riverside",
    both sharing one untagged file pool) would otherwise resolve to
    whichever building's file happens to sort first, regardless of which
    one was actually asked for; then across the full file list by keyword
    alone - so this still works when a property uses neither convention.

    files: list of UploadedFile (or None/empty) - the FULL shared pool.
    slot_rules: {slot_label: [(must_contain_all, must_not_contain_any), ...]}
                rules tried in order, first match wins; substrings matched
                case-insensitively against the filename.
    Returns {slot_label: UploadedFile or None}.
    """
    files = files or []
    by_name = {f.name: f for f in files}
    names = list(by_name.keys())
    assignment = {}
    used = set()

    tagged = [f for f in files if bucket_number is not None and has_bucket_tag(f.name.lower(), bucket_number)]
    hinted = [f for f in files if building_hint and building_hint.lower() in f.name.lower()]

    def matches_any_rule(fname_lower):
        for rule_groups in slot_rules.values():
            for must, must_not in rule_groups:
                if all(s.lower() in fname_lower for s in must) and not any(s.lower() in fname_lower for s in must_not):
                    return True
        return False

    # With every bucket's files dropped into one shared pool, a plain "every
    # uploaded file" dropdown gets long and hard to scan. Default each
    # dropdown to just the files that look relevant to THIS bucket, with an
    # opt-out checkbox for the rare true-manual-override case.
    relevant_names = [n for n in names if matches_any_rule(n.lower())]
    show_all = False
    if relevant_names and len(relevant_names) < len(names):
        show_all = st.checkbox(
            f"Show all {len(names)} uploaded files in these dropdowns (default: just the "
            f"{len(relevant_names)} that look relevant to this bucket)",
            key=f'{key_prefix}_show_all',
        )
    pool_names = names if (show_all or not relevant_names) else relevant_names

    for slot, rule_groups in slot_rules.items():
        guess = None
        for pool in (tagged, hinted, files):
            for f in pool:
                if f.name in used:
                    continue
                fname_lower = f.name.lower()
                for must, must_not in rule_groups:
                    if (all(s.lower() in fname_lower for s in must)
                            and not any(s.lower() in fname_lower for s in must_not)):
                        guess = f.name
                        break
                if guess:
                    break
            if guess:
                break
        options = ['(none)'] + pool_names
        if guess and guess not in options:
            options.append(guess)  # always selectable even if filtered out of the default view
        default_idx = options.index(guess) if guess in options else 0
        sel = st.selectbox(slot, options, index=default_idx, key=f'{key_prefix}_{slot}')
        if sel != '(none)':
            assignment[slot] = by_name[sel]
            used.add(sel)
        else:
            assignment[slot] = None
    return assignment


def silent_classify(files, slot_rules, bucket_number=None, building_hint=None):
    """Non-UI twin of classify_and_pick's auto-guess - identical (up to)
    three-pass (tagged, then building_hint substring, then full pool)
    matching, but returns {slot: file or None} directly with no dropdowns
    rendered. Used by 'Run Everything' to classify every bucket without
    visiting each tab first - always takes the auto-guess with no manual-
    override option (that's what the per-tab dropdowns remain for, if a
    guess is ever wrong). building_hint matters most here specifically,
    since Run Everything has no dropdown for a human to catch a wrong guess -
    see classify_and_pick's docstring for why it's needed at all."""
    files = files or []
    assignment = {}
    used = set()
    tagged = [f for f in files if bucket_number is not None and has_bucket_tag(f.name.lower(), bucket_number)]
    hinted = [f for f in files if building_hint and building_hint.lower() in f.name.lower()]
    for slot, rule_groups_for_slot in slot_rules.items():
        guess = None
        for pool in (tagged, hinted, files):
            for f in pool:
                if f.name in used:
                    continue
                fname_lower = f.name.lower()
                for must, must_not in rule_groups_for_slot:
                    if (all(s.lower() in fname_lower for s in must)
                            and not any(s.lower() in fname_lower for s in must_not)):
                        guess = f
                        break
                if guess:
                    break
            if guess:
                break
        assignment[slot] = guess
        if guess:
            used.add(guess.name)
    return assignment


def silent_classify_multi(files, rule_groups):
    """Non-UI twin of classify_multi's auto-guess - returns the matching
    files directly, no multiselect rendered."""
    def matches(fname_lower):
        for must, must_not in rule_groups:
            if all(s.lower() in fname_lower for s in must) and not any(s.lower() in fname_lower for s in must_not):
                return True
        return False
    return [f for f in (files or []) if matches(f.name.lower())]


def classify_multi(files, rule_groups, key, label):
    """Multiselect version of the auto-guess, for slots where a property may
    split one report into several files by category (e.g. Lexington Labs'
    Base Rent by use-type: Lab/Office/Misc, or a supplementary Rent Roll -
    Retail alongside the main roll) - a single dropdown can only ever pick
    one file, so this pre-selects every match and lets you add/remove."""
    files = files or []
    names = [f.name for f in files]
    by_name = {f.name: f for f in files}

    def matches(fname_lower):
        for must, must_not in rule_groups:
            if all(s.lower() in fname_lower for s in must) and not any(s.lower() in fname_lower for s in must_not):
                return True
        return False

    guesses = [n for n in names if matches(n.lower())]
    sel = st.multiselect(label, names, default=guesses, key=key)
    return [by_name[n] for n in sel]


def extract_tags(files):
    """Every distinct 'B<n>' filename tag present in this pool (e.g. Lexington
    Labs tags each of its 5 buildings' Bucket 1/6 files this way), sorted
    numerically as strings. Empty if the property doesn't use this convention."""
    tags = set()
    for f in files or []:
        for m in re.finditer(r'(?<![a-z0-9])b0*(\d+)(?![a-z0-9])', f.name.lower()):
            tags.add(m.group(1))
    return sorted(tags, key=int)


def classify_for_tag(files, slot_rules, tag):
    """Fully automatic version of classify_and_pick (no dropdown, so it's safe
    to call once per building in a loop) restricted to ONE building tag: for
    each slot, prefer a file carrying that exact 'B{tag}', falling back to
    untagged files only (never another building's tagged file) so a shared,
    portfolio-wide file - like a combined Summary - still gets picked up.
    Returns {slot_label: UploadedFile or None}."""
    files = files or []
    any_tag_re = re.compile(r'(?<![a-z0-9])b0*\d+(?![a-z0-9])')
    tagged = [f for f in files if has_bucket_tag(f.name.lower(), tag)]
    untagged = [f for f in files if not any_tag_re.search(f.name.lower())]
    assignment = {}
    used = set()
    for slot, rule_groups in slot_rules.items():
        pick = None
        for pool in (tagged, untagged):
            for f in pool:
                if f.name in used:
                    continue
                fname_lower = f.name.lower()
                for must, must_not in rule_groups:
                    if (all(s.lower() in fname_lower for s in must)
                            and not any(s.lower() in fname_lower for s in must_not)):
                        pick = f
                        break
                if pick:
                    break
            if pick:
                break
        assignment[slot] = pick
        if pick:
            used.add(pick.name)
    return assignment


def batch_runner(all_files, slot_rules, key_prefix, bucket_num, run_fn, store_extra=None):
    """Shown only when the file pool has more than one 'B<n>' building tag
    (e.g. Lexington Labs' 5 buildings, each with its own Detail/Monthly pair
    under one bucket) - the single dropdown above only ever picks one file per
    slot, so there's no way to analyze all of them without this. Runs the
    bucket's parser once per detected tag automatically (no manual dropdown -
    tag + keyword matching only) and stores each building's results under its
    own name, exactly like a normal single run.

    run_fn: (picked_dict, building_name, tag) -> results dict, e.g.
            lambda p, bname, tag: kardin_parser.run(p['Summary'], ...) - building_name is
            this tag's name (for cross-bucket lookups keyed by name, e.g. bucket 6
            needing bucket 1's results for the SAME building); tag is the raw 'B<n>'
            number, for callers that need to look up a property config entry.
    store_extra: optional (picked_dict) -> dict merged into the stored bucket_results entry
                 (e.g. {'detail_pdf': p['Budget Analysis Detail']})
    """
    tags = extract_tags(all_files)
    if len(tags) <= 1:
        return
    st.divider()
    st.subheader("Batch mode - multiple buildings detected")
    st.caption(f"{len(tags)} building tags found in the uploaded files: {', '.join('B' + t for t in tags)}. "
               "The dropdowns above only select one file per slot; this runs the bucket once per tag "
               "instead, fully automatically, and stores each as its own building.")
    property_cfg = st.session_state.get('property_cfg')
    names = {}
    cols = st.columns(len(tags))
    for col, t in zip(cols, tags):
        with col:
            cc = config_loader.cost_center_for_tag(property_cfg, t)
            default_name = cc['name'] if cc else f'B{t}'
            names[t] = st.text_input(f"Name for B{t}", value=default_name, key=f'{key_prefix}_batch_name_{t}')

    if st.button(f"Run Analysis for all {len(tags)} buildings", key=f'{key_prefix}_batch_run'):
        for t in tags:
            picked = classify_for_tag(all_files, slot_rules, t)
            bname = names[t]
            missing = [slot for slot, f in picked.items() if not f]
            if missing:
                st.warning(f"B{t} ('{bname}') couldn't auto-match: {', '.join(missing)} - "
                           "running anyway with what's available; each missing file is flagged below.")
            try:
                results = run_fn(picked, bname, t)
                entry = {'results': results}
                if store_extra:
                    entry.update(store_extra(picked))
                st.session_state.bucket_results[(bucket_num, bname)] = entry
                with st.expander(f"B{t} - '{bname}': {len(results['findings'])} finding(s)"):
                    show_checklist(results.get('checklist'))
                    show_stats(results['stats'])
                    show_findings(results['findings'])
                    if bucket_num == 1:
                        cc = config_loader.cost_center_for_tag(property_cfg, t)
                        show_expense_categorization(bname, cost_center_code=cc['code'] if cc else None)
                        show_assumptions_check(results['monthly_rows'], key_prefix=f'b1_batch_{t}')
                        show_interest_check(results['monthly_rows'], key_prefix=f'b1_batch_{t}')
            except Exception:
                st.error(f"B{t} ('{bname}') failed to parse:")
                st.code(traceback.format_exc())


st.header("Property")
_properties = config_loader.list_properties()
_prop_options = ['(none - manual building names only)'] + [p['name'] for p in _properties]
_prop_choice = st.selectbox(
    "Property config (optional)", _prop_options,
    help="Loads data/{property}/config.yaml - auto-fills batch mode's per-building names with the "
         "real cost-center name, and lets the parser flag a building's files if they're accidentally "
         "scoped to the wrong cost center. Not required - everything works with manual names too.",
)
st.session_state.property_slug = (
    next(p['slug'] for p in _properties if p['name'] == _prop_choice)
    if _prop_choice != _prop_options[0] else None
)
st.session_state.property_cfg = (
    config_loader.load_property_config(st.session_state.property_slug)
    if st.session_state.property_slug else None
)

with st.expander("Import / update cost centers from Kardin"):
    st.caption(
        "Upload Kardin's own 'Selected Cost Centers' report (rptSelectedCostCenters) - the "
        "authoritative list of every cost center defined for a property, straight from Kardin "
        "instead of hand-typed. Not every cost center is necessarily one of this tool's batch-mode "
        "buildings (a property can have cost centers for common areas, ground leases, etc. that "
        "were never part of the budget review) - check 'Include' only for the ones that are, and "
        "give each its 'B<n>' filename tag if it uses one."
    )
    roster_pdf = st.file_uploader("Selected Cost Centers report (.pdf)", type="pdf", key="cc_roster_upload")
    if roster_pdf:
        roster = kardin_parser.parse_cost_center_roster(roster_pdf)
        if not roster:
            st.warning("Couldn't find any cost center rows in that file - is it the rptSelectedCostCenters report?")
        else:
            st.caption(f"Found {len(roster)} cost center(s).")
            df = pd.DataFrame([{
                'Include': False, 'Tag (B<n>, optional)': '', 'Code': r['code'], 'Display Name': r['name'],
            } for r in roster])
            edited = st.data_editor(df, key='cc_roster_editor', hide_index=True, use_container_width=True)

            col1, col2 = st.columns(2)
            with col1:
                cc_prop_name = st.text_input("Property display name", key='cc_roster_propname')
            with col2:
                cc_slug = st.text_input("Folder slug (data/<slug>/config.yaml)", key='cc_roster_slug')

            if st.button("Generate config.yaml", key='cc_roster_generate'):
                included = edited[edited['Include']]
                if not cc_prop_name or not cc_slug or included.empty:
                    st.error("Need a property display name, a folder slug, and at least one included cost center.")
                else:
                    cost_centers = [{
                        'tag': (row['Tag (B<n>, optional)'] or '').strip() or None,
                        'code': row['Code'], 'name': row['Display Name'],
                    } for _, row in included.iterrows()]
                    yaml_text = config_loader.generate_config_yaml(cc_prop_name, cost_centers)
                    st.code(yaml_text, language='yaml')
                    st.download_button(
                        "Download config.yaml", data=yaml_text, file_name='config.yaml', mime='text/yaml',
                        key='cc_roster_download',
                    )
                    st.caption(f"Commit this as data/{cc_slug}/config.yaml on GitHub - Streamlit redeploys in "
                               "~1-2 minutes and it'll appear in the Property selector above.")

st.header("Building")
building = st.text_input("Building name (must match the tracker's tab name)", value=st.session_state.building)
st.session_state.building = building
st.caption("e.g. \"20 Riverside\" - exactly as it appears as a sheet tab in your Comment Tracker. "
           "Re-run a bucket if you change this and want its cross-checks against other buckets to line up.")

st.header("Files")
all_files = st.file_uploader(
    "Drop every file for this budget draft here - all buckets, both buildings if applicable, all at once. "
    "Each tab below auto-picks its own files from this pool (using a 'B1'/'B2'/... tag in the filename if "
    "present, otherwise by keyword) - override any guess with its dropdown.",
    type=["pdf", "xlsx"], accept_multiple_files=True, key="all_files",
)
if all_files:
    st.caption(f"{len(all_files)} file(s) uploaded.")
    with st.expander("Show uploaded filenames"):
        for f in all_files:
            st.write(f.name)


def run_everything(all_files, building):
    """Silently runs all 7 buckets in dependency order (1 before 4/5/6, 2
    before 7) using the SAME auto-classification every tab already uses -
    batch mode automatically for Bucket 1/6 when B<n> tags are detected,
    single-run otherwise, defaulting to ALL_BUILDINGS when no Building name
    is typed. Stores into st.session_state.bucket_results exactly like
    clicking each tab's own Run Analysis button, so every existing display
    (checklist, findings, categorization, etc.) works unchanged if the user
    later opens a tab. The optional assumption-file cross-checks (GRP Budget
    Assumptions, Leasing Assumptions, etc.) are NOT run here - those need an
    uploaded/committed file picked per-tab and stay exactly as they are.
    Returns [{'bucket', 'building', 'label', 'must_fix', 'for_discussion',
    'error'}] - one row per bucket run (one per batch building for 1/6)."""
    run_building = building if building else ALL_BUILDINGS
    summary = []

    def _tally(bucket_num, bname, label, findings):
        must_fix = sum(1 for f in findings if f.get('Priority') == 'Must Fix')
        summary.append({'bucket': bucket_num, 'building': bname, 'label': label,
                        'must_fix': must_fix, 'for_discussion': len(findings) - must_fix, 'error': None})

    property_cfg = st.session_state.get('property_cfg')

    # --- Bucket 1 ---
    b1_slot_rules = {
        'Budget Analysis Summary': [(['summary'], [])],
        'Budget Analysis Detail': [(['detail'], MONTHLY_TOKENS)],
        'Monthly Budget Detail': [([t], []) for t in MONTHLY_TOKENS],
    }
    tags = extract_tags(all_files)
    if len(tags) > 1:
        for t in tags:
            picked = classify_for_tag(all_files, b1_slot_rules, t)
            cc = config_loader.cost_center_for_tag(property_cfg, t)
            bname = cc['name'] if cc else f'B{t}'
            try:
                results = kardin_parser.run(
                    picked['Budget Analysis Summary'], picked['Budget Analysis Detail'],
                    picked['Monthly Budget Detail'], expected_cost_center=cc['code'] if cc else None)
                st.session_state.bucket_results[(1, bname)] = {
                    'results': results, 'detail_pdf': picked['Budget Analysis Detail']}
                _tally(1, bname, '1. Bgt & Fcst Summaries', results['findings'])
            except Exception:
                summary.append({'bucket': 1, 'building': bname, 'label': '1. Bgt & Fcst Summaries',
                                'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})
    else:
        picked = silent_classify(all_files, b1_slot_rules, bucket_number=1, building_hint=building)
        try:
            results = kardin_parser.run(picked['Budget Analysis Summary'], picked['Budget Analysis Detail'],
                                        picked['Monthly Budget Detail'])
            st.session_state.bucket_results[(1, building)] = {
                'results': results, 'detail_pdf': picked['Budget Analysis Detail']}
            _tally(1, building, '1. Bgt & Fcst Summaries', results['findings'])
        except Exception:
            summary.append({'bucket': 1, 'building': building, 'label': '1. Bgt & Fcst Summaries',
                            'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    # --- Bucket 2 ---
    picked = silent_classify(all_files, {
        'Free Rent': [(['free', 'rent'], [])],
        'Occupancy Summary': [(['occupancy'], [])],
        'Stacking Plan': [(['stacking'], [])],
    }, bucket_number=2, building_hint=building)
    rent_lab_pdfs = silent_classify_multi(all_files, [
        (['rent-lab'], ['free']), (['rent-office'], ['free']), (['rent-misc'], ['free']), (['base rent'], ['free'])])
    rent_roll_pdfs = silent_classify_multi(all_files, [(['rent roll'], [])])
    try:
        results = leasing_parser.run(picked['Free Rent'], rent_lab_pdfs, picked['Occupancy Summary'],
                                     rent_roll_pdfs, picked['Stacking Plan'])
        st.session_state.bucket_results[(2, run_building)] = {'results': results}
        _tally(2, run_building, '2. Leasing & Rent', results['findings'])
    except Exception:
        summary.append({'bucket': 2, 'building': run_building, 'label': '2. Leasing & Rent',
                        'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    # --- Bucket 3 ---
    picked = silent_classify(all_files, {
        'Recovery Calc Est': [(['calc est'], [])],
        'Recovery Monthly': [(['recovery', t], []) for t in MONTHLY_TOKENS],
        'Gross Up Schedule': [(['gross up'], [])],
    }, bucket_number=3, building_hint=building)
    xlsx_files = [f for f in (all_files or []) if f.name.lower().endswith('.xlsx')]
    picked_xlsx = silent_classify(xlsx_files, {'Fixed Factor Calcs (.xlsx)': [([], [])]}, bucket_number=3,
                                  building_hint=building)
    try:
        results = recoveries_parser.run(picked['Recovery Calc Est'], picked['Recovery Monthly'],
                                        picked['Gross Up Schedule'], picked_xlsx['Fixed Factor Calcs (.xlsx)'])
        st.session_state.bucket_results[(3, run_building)] = {'results': results}
        _tally(3, run_building, '3. Recoveries', results['findings'])
    except Exception:
        summary.append({'bucket': 3, 'building': run_building, 'label': '3. Recoveries',
                        'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    # --- Bucket 4 (cross-checks bucket 1) ---
    picked = silent_classify(all_files, {
        'Expense Detail': [(['expense detail'], [])],
        'Mgmt Fee Calc - File 1': [(['mgmt fee'], [])],
        'Mgmt Fee Calc - File 2 (optional)': [(['mgmt fee'], [])],
    }, bucket_number=4, building_hint=building)
    b1_entry = lookup_bucket(1, building)
    try:
        expense_detail_pdf = picked['Expense Detail']
        mgmt_fee_a_pdf = picked['Mgmt Fee Calc - File 1']
        mgmt_fee_b_pdf = picked['Mgmt Fee Calc - File 2 (optional)']
        bucket1_rows = b1_entry['results']['detail_rows'] if b1_entry else None
        results = expense_parser.run(
            expense_detail_pdf, mgmt_fee_a_pdf, mgmt_fee_b_pdf, bucket1_west20_detail_rows=bucket1_rows,
            mgmt_fee_20r_name=mgmt_fee_a_pdf.name if mgmt_fee_a_pdf else 'Mgmt Fee Calc #1',
            mgmt_fee_1r_name=mgmt_fee_b_pdf.name if mgmt_fee_b_pdf else 'Mgmt Fee Calc #2',
        )
        st.session_state.bucket_results[(4, run_building)] = {
            'results': results, 'expense_detail_name': expense_detail_pdf.name if expense_detail_pdf else 'Expense Detail'}
        _tally(4, run_building, '4. Expense Back-up', results['findings'])
    except Exception:
        summary.append({'bucket': 4, 'building': run_building, 'label': '4. Expense Back-up',
                        'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    # --- Bucket 5 (cross-checks bucket 1) ---
    picked = silent_classify(all_files, {
        'Capex': [(['capex'], [])],
        'TIs': [(['tis'], []), (['tenant improvement'], [])],
        'LCs': [(['lcs'], ['calc']), (['leasing commission'], [])],
    }, bucket_number=5, building_hint=building)
    try:
        capex_line_rows, capex_totals_rows = expense_parser.parse_expense_detail(picked['Capex'])
        bucket1_rows = b1_entry['results']['detail_rows'] if b1_entry else None
        results = capex_parser.run(picked['TIs'], picked['LCs'], capex_line_rows=capex_line_rows,
                                   capex_totals_rows=capex_totals_rows, bucket1_west20_detail_rows=bucket1_rows,
                                   capex_pdf_missing=(picked['Capex'] is None))
        st.session_state.bucket_results[(5, run_building)] = {'results': results, 'capex_line_rows': capex_line_rows}
        _tally(5, run_building, '5. CapEx Back-up', results['findings'])
    except Exception:
        summary.append({'bucket': 5, 'building': run_building, 'label': '5. CapEx Back-up',
                        'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    # --- Bucket 6 (cross-checks bucket 1) ---
    b6_slot_rules = {
        '2026B v 2026F Detail': [(['detail'], MONTHLY_TOKENS)],
        '2026F Monthly Detail': [([t], []) for t in MONTHLY_TOKENS],
    }
    b6_files = exclude_3way_reports(all_files)
    if len(tags) > 1:
        for t in tags:
            picked = classify_for_tag(b6_files, b6_slot_rules, t)
            cc = config_loader.cost_center_for_tag(property_cfg, t)
            bname = cc['name'] if cc else f'B{t}'
            b1e = st.session_state.bucket_results.get((1, bname))
            try:
                results = forecast_parser.run(
                    picked['2026B v 2026F Detail'], picked['2026F Monthly Detail'],
                    bucket1_detail_rows=b1e['results']['detail_rows'] if b1e else None,
                    bucket1_detail_pdf=b1e['detail_pdf'] if b1e else None)
                st.session_state.bucket_results[(6, bname)] = {'results': results}
                _tally(6, bname, '6. Forecast Back-up', results['findings'])
            except Exception:
                summary.append({'bucket': 6, 'building': bname, 'label': '6. Forecast Back-up',
                                'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})
    else:
        picked = silent_classify(b6_files, b6_slot_rules, bucket_number=6, building_hint=building)
        try:
            results = forecast_parser.run(
                picked['2026B v 2026F Detail'], picked['2026F Monthly Detail'],
                bucket1_detail_rows=b1_entry['results']['detail_rows'] if b1_entry else None,
                bucket1_detail_pdf=b1_entry['detail_pdf'] if b1_entry else None)
            st.session_state.bucket_results[(6, building)] = {'results': results}
            _tally(6, building, '6. Forecast Back-up', results['findings'])
        except Exception:
            summary.append({'bucket': 6, 'building': building, 'label': '6. Forecast Back-up',
                            'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    # --- Bucket 7 (cross-checks bucket 2) ---
    picked = silent_classify(all_files, {
        'Lease Expiration Schedule': [(['lease expiration'], []), (['lease', 'exp'], [])],
    }, bucket_number=7, building_hint=building)
    b2_entry = lookup_bucket(2, building)
    try:
        results = xtra_parser.run(picked['Lease Expiration Schedule'],
                                  occupancy_rows=b2_entry['results']['occupancy_rows'] if b2_entry else None)
        st.session_state.bucket_results[(7, run_building)] = {'results': results}
        _tally(7, run_building, '7. Xtra rpts', results['findings'])
    except Exception:
        summary.append({'bucket': 7, 'building': run_building, 'label': '7. Xtra rpts',
                        'must_fix': None, 'for_discussion': None, 'error': str(traceback.format_exc())})

    return summary


with st.expander("⚡ Run Everything", expanded=bool(all_files)):
    st.caption("Runs all 7 buckets at once using auto-classification only (no manual dropdown overrides) - "
               "Bucket 1/6 batch automatically across every detected building, the rest use the Building name "
               "above (or run portfolio-wide if it's blank). Results land in st.session_state exactly like "
               "clicking each tab's own Run Analysis, so every tab's display still works normally afterward. "
               "The optional assumption-file cross-checks (GRP Budget Assumptions, etc.) aren't included here - "
               "open the relevant tab for those.")
    if st.button("Run all 7 buckets now", type="primary", key="run_everything_btn", disabled=not all_files):
        with st.spinner("Running all 7 buckets..."):
            st.session_state.run_everything_summary = run_everything(all_files, building)

    _summary = st.session_state.get('run_everything_summary')
    if _summary:
        total_must_fix = sum(r['must_fix'] or 0 for r in _summary)
        total_discuss = sum(r['for_discussion'] or 0 for r in _summary)
        total_errors = sum(1 for r in _summary if r['error'])
        col1, col2, col3 = st.columns(3)
        col1.metric("Must Fix (all buckets)", total_must_fix)
        col2.metric("For Discussion (all buckets)", total_discuss)
        col3.metric("Buckets that failed to parse", total_errors)
        summary_df = pd.DataFrame([{
            'Bucket': r['label'], 'Building': r['building'],
            'Must Fix': r['must_fix'] if r['error'] is None else '—',
            'For Discussion': r['for_discussion'] if r['error'] is None else '—',
            'Status': 'ERROR - see traceback below' if r['error'] else 'OK',
        } for r in _summary])
        st.dataframe(summary_df, use_container_width=True, hide_index=True)
        for r in _summary:
            if r['error']:
                with st.expander(f"⚠️ {r['label']} ({r['building']}) failed - traceback"):
                    st.code(r['error'])
        st.caption("Open any tab below to see that bucket's full findings, checklist, and optional "
                   "cross-checks - this summary is just the headline count.")

tabs = st.tabs([
    "1. Bgt & Fcst Summaries", "2. Leasing & Rent", "3. Recoveries", "4. Expense Back-up",
    "5. CapEx Back-up", "6. Forecast Back-up", "7. Xtra rpts", "Merge to Tracker",
])

# ------------------------------------------------------------- 1. Bgt & Fcst Summaries
with tabs[0]:
    st.caption("Budget Analysis Summary / Detail / Monthly Detail")
    picked = classify_and_pick(all_files, {
        'Budget Analysis Summary': [(['summary'], [])],
        'Budget Analysis Detail': [(['detail'], MONTHLY_TOKENS)],
        'Monthly Budget Detail': [([t], []) for t in MONTHLY_TOKENS],
    }, 'b1', bucket_number=1, building_hint=building)
    summary_pdf = picked['Budget Analysis Summary']
    detail_pdf = picked['Budget Analysis Detail']
    monthly_pdf = picked['Monthly Budget Detail']

    if run_button("b1_run", {
        'Budget Analysis Summary': bool(summary_pdf), 'Budget Analysis Detail': bool(detail_pdf),
        'Monthly Budget Detail': bool(monthly_pdf), 'Building name': bool(building),
    }):
        try:
            results = kardin_parser.run(summary_pdf, detail_pdf, monthly_pdf)
            st.session_state.bucket_results[(1, building)] = {'results': results, 'detail_pdf': detail_pdf}
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((1, building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])
        _cc = config_loader.cost_center_for_name(st.session_state.get('property_cfg'), building)
        show_expense_categorization(building, cost_center_code=_cc['code'] if _cc else None)
        show_assumptions_check(entry['results']['monthly_rows'], key_prefix='b1_single')
        show_interest_check(entry['results']['monthly_rows'], key_prefix='b1_single')

    b1_slot_rules = {
        'Budget Analysis Summary': [(['summary'], [])],
        'Budget Analysis Detail': [(['detail'], MONTHLY_TOKENS)],
        'Monthly Budget Detail': [([t], []) for t in MONTHLY_TOKENS],
    }
    def _b1_run(picked, bname, tag):
        cc = config_loader.cost_center_for_tag(st.session_state.get('property_cfg'), tag)
        return kardin_parser.run(
            picked['Budget Analysis Summary'], picked['Budget Analysis Detail'], picked['Monthly Budget Detail'],
            expected_cost_center=cc['code'] if cc else None,
        )

    batch_runner(
        all_files, b1_slot_rules, 'b1', 1, run_fn=_b1_run,
        store_extra=lambda p: {'detail_pdf': p['Budget Analysis Detail']},
    )

# ------------------------------------------------------------------- 2. Leasing & Rent
with tabs[1]:
    st.caption("Free Rent / Base Rent (Rent-Lab/Office/Misc-Mnthly) / Occupancy Summary / Rent Roll / Stacking Plan")
    picked = classify_and_pick(all_files, {
        'Free Rent': [(['free', 'rent'], [])],
        'Occupancy Summary': [(['occupancy'], [])],
        'Stacking Plan': [(['stacking'], [])],
    }, 'b2', bucket_number=2, building_hint=building)
    free_rent_pdf = picked['Free Rent']
    occupancy_pdf = picked['Occupancy Summary']
    stacking_pdf = picked['Stacking Plan']

    rent_lab_pdfs = classify_multi(
        all_files,
        [(['rent-lab'], ['free']), (['rent-office'], ['free']), (['rent-misc'], ['free']), (['base rent'], ['free'])],
        'b2_rent_lab_multi',
        "Base Rent - select every use-type file that applies (Lab / Office / Misc, etc.)",
    )
    rent_roll_pdfs = classify_multi(
        all_files, [(['rent roll'], [])], 'b2_rent_roll_multi',
        "Rent Roll - select every file that applies (main + any Retail/other supplement)",
    )

    eff_building, building_ok = building_selector('b2', building)

    if run_button("b2_run", {
        'Free Rent': bool(free_rent_pdf), 'Base Rent (at least one)': bool(rent_lab_pdfs),
        'Occupancy Summary': bool(occupancy_pdf), 'Rent Roll (at least one)': bool(rent_roll_pdfs),
        'Stacking Plan': bool(stacking_pdf), 'Building name': building_ok,
    }):
        try:
            results = leasing_parser.run(free_rent_pdf, rent_lab_pdfs, occupancy_pdf, rent_roll_pdfs, stacking_pdf)
            st.session_state.bucket_results[(2, eff_building)] = {'results': results}
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((2, eff_building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])
        show_leasing_assumptions_check(entry['results']['occupancy_rows'], key_prefix='b2_single')

# --------------------------------------------------------------------------- 3. Recoveries
with tabs[2]:
    st.caption("Recovery Calc Est / Recovery Monthly / Gross Up Schedule / Fixed Factor Calcs (.xlsx)")
    picked = classify_and_pick(all_files, {
        'Recovery Calc Est': [(['calc est'], [])],
        'Recovery Monthly': [(['recovery', t], []) for t in MONTHLY_TOKENS],
        'Gross Up Schedule': [(['gross up'], [])],
    }, 'b3', bucket_number=3, building_hint=building)
    calc_est_pdf = picked['Recovery Calc Est']
    recovery_monthly_pdf = picked['Recovery Monthly']
    gross_up_pdf = picked['Gross Up Schedule']

    xlsx_files = [f for f in (all_files or []) if f.name.lower().endswith('.xlsx')]
    picked_xlsx = classify_and_pick(xlsx_files, {'Fixed Factor Calcs (.xlsx)': [([], [])]}, 'b3x', bucket_number=3,
                                    building_hint=building)
    fixed_factor_xlsx = picked_xlsx['Fixed Factor Calcs (.xlsx)']

    eff_building, building_ok = building_selector('b3', building)

    if run_button("b3_run", {
        'Recovery Calc Est': bool(calc_est_pdf), 'Recovery Monthly': bool(recovery_monthly_pdf),
        'Gross Up Schedule': bool(gross_up_pdf), 'Fixed Factor Calcs (.xlsx)': bool(fixed_factor_xlsx),
        'Building name': building_ok,
    }):
        try:
            results = recoveries_parser.run(calc_est_pdf, recovery_monthly_pdf, gross_up_pdf, fixed_factor_xlsx)
            st.session_state.bucket_results[(3, eff_building)] = {'results': results}
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((3, eff_building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])

# ----------------------------------------------------------------------- 4. Expense Back-up
with tabs[3]:
    st.caption("Expense Detail / Mgmt Fee Calc x2. Cross-checks against bucket 1's Detail when available.")
    picked = classify_and_pick(all_files, {
        'Expense Detail': [(['expense detail'], [])],
        'Mgmt Fee Calc - File 1': [(['mgmt fee'], [])],
        'Mgmt Fee Calc - File 2 (optional)': [(['mgmt fee'], [])],
    }, 'b4', bucket_number=4, building_hint=building)
    expense_detail_pdf = picked['Expense Detail']
    mgmt_fee_a_pdf = picked['Mgmt Fee Calc - File 1']
    mgmt_fee_b_pdf = picked['Mgmt Fee Calc - File 2 (optional)']
    st.caption("Second Mgmt Fee Calc file is optional - Kardin dumps the same full-portfolio calc into "
               "every export regardless of which building was requested, so some PMs only send one.")

    b1_entry = lookup_bucket(1, building)
    if b1_entry:
        st.caption("✓ Bucket 1 data found for this building - GL tie-out will run.")
    else:
        st.caption("Bucket 1 not yet run for this building - GL tie-out will be skipped.")

    eff_building, building_ok = building_selector('b4', building)

    if run_button("b4_run", {
        'Expense Detail': bool(expense_detail_pdf), 'Mgmt Fee Calc - File 1': bool(mgmt_fee_a_pdf),
        'Building name': building_ok,
    }):
        try:
            bucket1_rows = b1_entry['results']['detail_rows'] if b1_entry else None
            results = expense_parser.run(
                expense_detail_pdf, mgmt_fee_a_pdf, mgmt_fee_b_pdf,
                bucket1_west20_detail_rows=bucket1_rows,
                mgmt_fee_20r_name=mgmt_fee_a_pdf.name if mgmt_fee_a_pdf else 'Mgmt Fee Calc #1',
                mgmt_fee_1r_name=mgmt_fee_b_pdf.name if mgmt_fee_b_pdf else 'Mgmt Fee Calc #2',
            )
            st.session_state.bucket_results[(4, eff_building)] = {
                'results': results,
                'expense_detail_name': expense_detail_pdf.name if expense_detail_pdf else 'Expense Detail',
            }
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((4, eff_building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])
        show_expense_yoy_comparison(entry['results']['line_rows'], key_prefix='b4_single')

# ------------------------------------------------------------------------ 5. CapEx Back-up
with tabs[4]:
    st.caption("Capex / TIs / LCs. Cross-checks against bucket 1's Detail when available.")
    picked = classify_and_pick(all_files, {
        'Capex': [(['capex'], [])],
        'TIs': [(['tis'], []), (['tenant improvement'], [])],
        'LCs': [(['lcs'], ['calc']), (['leasing commission'], [])],
    }, 'b5', bucket_number=5, building_hint=building)
    capex_pdf = picked['Capex']
    tis_pdf = picked['TIs']
    lcs_pdf = picked['LCs']

    b1_entry = lookup_bucket(1, building)
    if b1_entry:
        st.caption("✓ Bucket 1 data found for this building - GL tie-out will run.")
    else:
        st.caption("Bucket 1 not yet run for this building - GL tie-out will be skipped.")

    eff_building, building_ok = building_selector('b5', building)

    if run_button("b5_run", {
        'Capex': bool(capex_pdf), 'TIs': bool(tis_pdf), 'LCs': bool(lcs_pdf), 'Building name': building_ok,
    }):
        try:
            capex_line_rows, capex_totals_rows = expense_parser.parse_expense_detail(capex_pdf)
            bucket1_rows = b1_entry['results']['detail_rows'] if b1_entry else None
            results = capex_parser.run(tis_pdf, lcs_pdf, capex_line_rows=capex_line_rows,
                                       capex_totals_rows=capex_totals_rows,
                                       bucket1_west20_detail_rows=bucket1_rows,
                                       capex_pdf_missing=(capex_pdf is None))
            st.session_state.bucket_results[(5, eff_building)] = {
                'results': results, 'capex_line_rows': capex_line_rows,
            }
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((5, eff_building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])
        show_capex_plan_check(entry.get('capex_line_rows') or [], key_prefix='b5_single')

# -------------------------------------------------------------------- 6. Forecast Back-up
with tabs[5]:
    st.caption("2026B v 2026F Detail / 2026F Monthly Detail. Cross-checks against bucket 1's Detail when available.")
    picked = classify_and_pick(exclude_3way_reports(all_files), {
        '2026B v 2026F Detail': [(['detail'], MONTHLY_TOKENS)],
        '2026F Monthly Detail': [([t], []) for t in MONTHLY_TOKENS],
    }, 'b6', bucket_number=6, building_hint=building)
    fc_detail_pdf = picked['2026B v 2026F Detail']
    fc_monthly_pdf = picked['2026F Monthly Detail']

    b1_entry = lookup_bucket(1, building)
    if b1_entry:
        st.caption("✓ Bucket 1 data found for this building - reforecast drift check will run.")
    else:
        st.caption("Bucket 1 not yet run for this building - reforecast drift check will be skipped.")

    if run_button("b6_run", {
        '2026B v 2026F Detail': bool(fc_detail_pdf), '2026F Monthly Detail': bool(fc_monthly_pdf),
        'Building name': bool(building),
    }):
        try:
            bucket1_rows = b1_entry['results']['detail_rows'] if b1_entry else None
            bucket1_pdf = b1_entry['detail_pdf'] if b1_entry else None
            results = forecast_parser.run(fc_detail_pdf, fc_monthly_pdf,
                                          bucket1_detail_rows=bucket1_rows, bucket1_detail_pdf=bucket1_pdf)
            st.session_state.bucket_results[(6, building)] = {'results': results}
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((6, building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])

    b6_slot_rules = {
        '2026B v 2026F Detail': [(['detail'], MONTHLY_TOKENS)],
        '2026F Monthly Detail': [([t], []) for t in MONTHLY_TOKENS],
    }

    def _b6_run(picked, bname, tag):
        b1e = st.session_state.bucket_results.get((1, bname))
        return forecast_parser.run(
            picked['2026B v 2026F Detail'], picked['2026F Monthly Detail'],
            bucket1_detail_rows=b1e['results']['detail_rows'] if b1e else None,
            bucket1_detail_pdf=b1e['detail_pdf'] if b1e else None,
        )

    st.caption("Batch mode below looks up each building's bucket-1 results (by the same name typed there) "
               "for the reforecast drift check - run bucket 1's batch first if you want that check included.")
    batch_runner(exclude_3way_reports(all_files), b6_slot_rules, 'b6', 6, run_fn=_b6_run)

# --------------------------------------------------------------------------- 7. Xtra rpts
with tabs[6]:
    st.caption("Lease Expiration Schedule. Cross-checks against bucket 2's Occupancy Summary when available.")
    picked = classify_and_pick(all_files, {
        'Lease Expiration Schedule': [(['lease expiration'], []), (['lease', 'exp'], [])],
    }, 'b7', bucket_number=7, building_hint=building)
    lease_exp_pdf = picked['Lease Expiration Schedule']

    b2_entry = lookup_bucket(2, building)
    if b2_entry:
        st.caption("✓ Bucket 2 data found for this building - Occupancy Summary tie-out will run.")
    else:
        st.caption("Bucket 2 not yet run for this building - Occupancy Summary tie-out will be skipped.")

    eff_building, building_ok = building_selector('b7', building)

    if run_button("b7_run", {'Lease Expiration Schedule': bool(lease_exp_pdf), 'Building name': building_ok}):
        try:
            occupancy_rows = b2_entry['results']['occupancy_rows'] if b2_entry else None
            results = xtra_parser.run(lease_exp_pdf, occupancy_rows=occupancy_rows)
            st.session_state.bucket_results[(7, eff_building)] = {'results': results}
            st.success(f"Parsed - {len(results['findings'])} finding(s).")
        except Exception:
            report_error()

    entry = st.session_state.bucket_results.get((7, eff_building))
    if entry:
        show_checklist(entry['results'].get('checklist'))
        show_stats(entry['results']['stats'])
        show_findings(entry['results']['findings'])

# --------------------------------------------------------------------------- Merge to Tracker
with tabs[7]:
    st.caption("Combines findings from every bucket you've analyzed this session for this building.")

    combined_findings = []
    bucket_labels = {
        1: "1. Bgt & Fcst Summaries", 2: "2. Leasing & Rent", 3: "3. Recoveries",
        4: "4. Expense Back-up", 5: "5. CapEx Back-up", 6: "6. Forecast Back-up", 7: "7. Xtra rpts",
    }
    present = []
    for n in range(1, 8):
        # Include both this specific building's own run AND any portfolio-
        # wide "(All Buildings)" run of the same bucket (see building_selector) -
        # a bucket can be one, the other, or (rarely) both.
        e_specific = st.session_state.bucket_results.get((n, building))
        e_all = st.session_state.bucket_results.get((n, ALL_BUILDINGS))
        if e_specific:
            combined_findings += e_specific['results']['findings']
        if e_all:
            combined_findings += e_all['results']['findings']
        if e_specific or e_all:
            present.append(bucket_labels[n])
    missing = [v for k, v in bucket_labels.items() if bucket_labels[k] not in present]

    if present:
        st.write(f"**Included:** {', '.join(present)}")
    if missing:
        st.write(f"**Not yet run for '{building}':** {', '.join(missing)}")

    st.metric("Total findings to merge", len(combined_findings))

    tracker_file = st.file_uploader("Upload the current GRP Budget Comment Tracker (.xlsx)", type="xlsx",
                                     key="tracker_upload")

    if tracker_file and combined_findings:
        try:
            wb_probe = openpyxl.load_workbook(tracker_file, data_only=True, read_only=True)
            sheet_names = wb_probe.sheetnames
            tracker_file.seek(0)
        except Exception:
            sheet_names = []
            st.error("Couldn't read that file as an Excel workbook.")
            st.code(traceback.format_exc())

        default_idx = sheet_names.index(building) if building in sheet_names else 0
        if sheet_names:
            sheet_name = st.selectbox("Which tab is this building?", sheet_names, index=default_idx)
            comment_by = st.selectbox(
                "Comment By (attributed reviewer)",
                ["Ryan Walsh", "Lauren Sullivan", "Natasha Parker", "Justin Bentayou"],
                index=0,
            )

            if st.button("Merge & prepare download"):
                try:
                    tracker_file.seek(0)
                    buf, n_inserted, warnings = tracker_writer.merge_findings_into_tracker(
                        tracker_file, sheet_name, combined_findings, comment_by=comment_by
                    )
                    for w in warnings:
                        st.warning(w)
                    st.success(f"Inserted {n_inserted} row(s) into '{sheet_name}'.")
                    out_name = f"GRP Budget Comment Tracker (with QC - {date.today().isoformat()}).xlsx"
                    st.download_button(
                        "Download updated tracker",
                        data=buf,
                        file_name=out_name,
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                except tracker_writer.TrackerStructureError as e:
                    st.error(f"Tracker structure issue: {e}")
                except Exception:
                    st.error("Merge failed. Full traceback below.")
                    st.code(traceback.format_exc())
    elif not combined_findings and tracker_file:
        st.info("No findings to merge yet - run at least one bucket above first.")
